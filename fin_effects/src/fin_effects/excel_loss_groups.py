"""Prod-парсер выгрузок графового сервиса: широкие Excel с колонкой «Убыток».

Правила:

* лист берётся, если в шапке есть колонка «Убыток» (после нормализации имени);
* в группу попадают строки, где **хотя бы одна** ячейка имеет цвет заливки
  или цвет шрифта (любой, отличный от «пустого» / авто / чёрного по умолчанию);
* **1 файл = 1 группа** (``group_id`` = stem имени файла); листы склеиваются;
* из цветных строк: номер убытка + все пары ФИО+ДР по эвристикам заголовков.

Чтение стилей — через ``openpyxl`` (``pd.read_excel`` цвет не отдаёт).
"""

from __future__ import annotations

import logging
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.worksheet.worksheet import Worksheet

from .textnorm import format_person_key, normalize_text

LOGGER = logging.getLogger(__name__)

#: Сколько первых строк листа просматриваем в поисках шапки.
HEADER_SCAN_ROWS = 50

#: Итоговые колонки prod-выгрузки участников.
OUTPUT_COLUMNS = (
    "full_name",
    "birth_date",
    "group_id",
    "person_key",
    "loss_number",
    "source_file",
    "source_sheet",
    "row_index",
)

_WHITESPACE_RE = re.compile(r"[\s\u00a0]+")
_DOT_SPACE_RE = re.compile(r"\s*\.\s*")

#: Токены в заголовке, которые делают колонку «не ФИО» (ТС, адреса, …).
_NAME_EXCLUDE_TOKENS = (
    "марка",
    "модель",
    "гос номер",
    "госномер",
    "регистрац",
    "идентификатор",
    "vin",
    "полис",
    "эмитент",
    "адрес",
    "телефон",
    "паспорт",
    "объект",
    "год выпуск",
    "выплат",
    "сумма",
    "решени",
    "стадия",
    "филиал",
    "место",
    "гибдд",
    "коммент",
    "проверка",
    "основан",
    "экспертиз",
)

#: Подстроки заголовка, указывающие на ФИО / наименование лица.
_NAME_INCLUDE_TOKENS = (
    "полн наименование",
    "полное наименование",
    "лицо в именительном",
    "наименование",
    "заявитель",
    "фамилия",
    "получатель возмещения",
)

#: Роли, у которых колонка может быть просто «Виновник» / «Потерпевший» без «наименование».
_SIMPLE_NAME_ROLES = (
    "виновник",
    "потерпевший",
    "страхователь",
    "заявитель",
    "владелец тс",
    "получатель возмещения",
)

#: Подстроки заголовка даты рождения.
_BIRTH_TOKENS = (
    "дата рожд",
    "датарождения",
    "д р",
    "др ",
    " др",
)

#: RGB/theme, которые считаем «без цвета» для шрифта (чёрный / авто).
_DEFAULT_FONT_RGB = frozenset(
    {
        "00000000",
        "000000",
        "FF000000",
        "01000000",
    }
)


@dataclass(frozen=True)
class ProdSheetStats:
    """Статистика разбора одного листа."""

    sheets_with_ubytok: int = 0
    colored_rows: int = 0
    loss_numbers: int = 0
    person_keys: int = 0
    skipped_no_header: int = 0


def normalize_header(value: object) -> str:
    """Нормализует имя колонки для поиска (регистр, пробелы, точки).

    Args:
        value: сырое значение ячейки шапки.

    Returns:
        Строка в нижнем регистре без лишних пробелов.
    """
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    text = str(value).replace("\u00a0", " ")
    text = _WHITESPACE_RE.sub(" ", text)
    text = _DOT_SPACE_RE.sub(".", text)
    text = text.strip().casefold()
    text = _WHITESPACE_RE.sub(" ", text).strip()
    return text


def _role_prefix(normalized: str) -> str:
    """Префикс роли до последнего сегмента (для пары ФИО+ДР)."""
    if "." in normalized:
        return normalized.rsplit(".", 1)[0].strip()
    # «Виновник Дата рождения» / «Виновник Полн наименование»
    for token in (" дата рожд", " датарожд", " полн", " полное", " лицо"):
        idx = normalized.find(token)
        if idx > 0:
            return normalized[:idx].strip()
    return normalized


def _is_birth_header(normalized: str) -> bool:
    if not normalized:
        return False
    if any(ex in normalized for ex in ("событ", "обращен", "заявлен", "создан", "закрыт")):
        # «Дата события» и т.п. — не ДР
        if "рожд" not in normalized and "д р" not in normalized:
            return False
    return any(token in normalized for token in _BIRTH_TOKENS)


def _is_name_header(normalized: str) -> bool:
    if not normalized or _is_birth_header(normalized):
        return False
    if any(ex in normalized for ex in _NAME_EXCLUDE_TOKENS):
        return False
    if any(token in normalized for token in _NAME_INCLUDE_TOKENS):
        return True
    # Простая колонка роли без «наименование»: «Виновник», «Потерпевший»
    for role in _SIMPLE_NAME_ROLES:
        if normalized == role or normalized.endswith("." + role) or normalized.endswith(" " + role):
            return True
        if normalized.startswith(role + " ") and "дата" not in normalized:
            # «Виновник П/Ю/н/наименование» уже покрыто include; короткие имена
            if "номер" in normalized or "полис" in normalized:
                return False
    return False


def _is_loss_number_header(normalized: str) -> bool:
    """Колонка номера убытка: ровно «убыток», не «убыток.код…»."""
    return normalized == "убыток"


def _header_has_ubytok(headers: Sequence[str]) -> bool:
    return any(h == "убыток" or h.startswith("убыток.") or h.startswith("убыток ") for h in headers)


@dataclass(frozen=True)
class NameBirthSpec:
    """Как собрать ФИО+ДР из колонок строки."""

    name_cols: tuple[int, ...]
    birth_col: int | None


def build_name_birth_specs(headers: Sequence[str]) -> list[NameBirthSpec]:
    """Строит спецификации извлечения ФИО+ДР для листа."""
    specs: list[NameBirthSpec] = []
    name_cols = [i for i, h in enumerate(headers) if _is_name_header(h)]
    birth_cols = [i for i, h in enumerate(headers) if _is_birth_header(h)]
    used_birth: set[int] = set()

    for name_idx in name_cols:
        name_h = headers[name_idx]
        prefix = _role_prefix(name_h)
        best: int | None = None
        best_score = -1
        for birth_idx in birth_cols:
            if birth_idx in used_birth:
                continue
            birth_h = headers[birth_idx]
            birth_prefix = _role_prefix(birth_h)
            score = 0
            if prefix and (
                prefix == birth_prefix
                or birth_prefix.startswith(prefix)
                or prefix.startswith(birth_prefix)
            ):
                score = 100 - abs(birth_idx - name_idx)
            elif abs(birth_idx - name_idx) <= 2:
                score = 50 - abs(birth_idx - name_idx)
            if score > best_score:
                best_score = score
                best = birth_idx
        if best is not None and best_score >= 0:
            used_birth.add(best)
            specs.append(NameBirthSpec(name_cols=(name_idx,), birth_col=best))
        else:
            specs.append(NameBirthSpec(name_cols=(name_idx,), birth_col=None))

    # Фамилия + Имя (+ Отчество)
    for i, h in enumerate(headers):
        if not (h.endswith("фамилия") or h == "фамилия"):
            continue
        prefix = _role_prefix(h)
        parts = [i]
        for suffix in ("имя", "отчество"):
            found = next(
                (
                    j
                    for j, hh in enumerate(headers)
                    if (hh.endswith(suffix) or hh == suffix) and _role_prefix(hh) == prefix
                ),
                None,
            )
            if found is not None:
                parts.append(found)
        if len(parts) < 2:
            continue
        birth_idx = next(
            (
                j
                for j in birth_cols
                if j not in used_birth and _role_prefix(headers[j]) == prefix
            ),
            None,
        )
        if birth_idx is not None:
            used_birth.add(birth_idx)
        specs.append(NameBirthSpec(name_cols=tuple(parts), birth_col=birth_idx))

    return specs


def _theme_or_rgb(color: Any) -> str | None:
    """Достаёт RGB/indexed-строку цвета openpyxl или ``None``."""
    if color is None or color.type is None:
        return None
    if color.type == "rgb" and color.rgb:
        rgb = str(color.rgb).upper()
        if rgb in {"00000000", "0"}:
            return None
        return rgb
    if color.type == "indexed" and color.indexed is not None:
        try:
            indexed = int(color.indexed)
        except (TypeError, ValueError):
            return None
        if indexed in (0, 1, 64, 65):  # auto / system
            return None
        if 0 <= indexed < len(COLOR_INDEX):
            return str(COLOR_INDEX[indexed]).upper()
        return f"INDEXED:{indexed}"
    if color.type == "theme":
        # theme без tint — часто «авто»; с tint — считается окрашенным
        if color.theme in (0, 1) and not color.tint:
            return None
        return f"THEME:{color.theme}:{color.tint}"
    return None


def cell_has_color(cell: Cell) -> bool:
    """Есть ли у ячейки значимая заливка или цвет шрифта."""
    fill = cell.fill
    if fill is not None and fill.fill_type not in (None, "none"):
        fg = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
        bg = getattr(fill, "bgColor", None) or getattr(fill, "end_color", None)
        for color in (fg, bg):
            rgb = _theme_or_rgb(color)
            if rgb and rgb not in _DEFAULT_FONT_RGB:
                # Белый/почти прозрачный часто = «нет заливки»
                if rgb.endswith("FFFFFF") or rgb in {"FFFFFFFF", "00FFFFFF"}:
                    continue
                return True

    font = cell.font
    if font is not None and font.color is not None:
        rgb = _theme_or_rgb(font.color)
        if rgb and rgb not in _DEFAULT_FONT_RGB and not rgb.endswith("000000"):
            return True
    return False


def row_has_color(worksheet: Worksheet, row: int, max_col: int) -> bool:
    """Есть ли в строке хотя бы одна окрашенная ячейка."""
    for col in range(1, max_col + 1):
        if cell_has_color(worksheet.cell(row=row, column=col)):
            return True
    return False


def _cell_value(cell: Cell) -> object:
    return cell.value


def _normalize_loss_number(value: object) -> str:
    """Номер убытка без хвостового ``.0`` у float."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    elif isinstance(value, int):
        pass
    else:
        text = str(value).strip()
        if re.fullmatch(r"\d+\.0+", text):
            text = text.split(".", 1)[0]
        value = text
    return normalize_text(value)


def find_header_row(
    worksheet: Worksheet,
    *,
    max_scan: int = HEADER_SCAN_ROWS,
) -> tuple[int, list[str], list[object]] | None:
    """Ищет строку шапки с колонкой «Убыток».

    Returns:
        ``(row_1based, normalized_headers, raw_headers)`` или ``None``.
    """
    max_row = min(worksheet.max_row or 0, max_scan)
    max_col = worksheet.max_column or 0
    if max_row < 1 or max_col < 1:
        return None

    for row in range(1, max_row + 1):
        raw = [_cell_value(worksheet.cell(row=row, column=col)) for col in range(1, max_col + 1)]
        headers = [normalize_header(value) for value in raw]
        if not _header_has_ubytok(headers):
            continue
        if not any(_is_name_header(h) or h.endswith("фамилия") for h in headers):
            # Лист только с «Убыток» без ФИО всё равно полезен (матч по номеру)
            if not any(_is_loss_number_header(h) for h in headers):
                continue
        return row, headers, raw
    return None


def find_loss_column(headers: Sequence[str]) -> int | None:
    """Индекс колонки номера убытка (0-based) или ``None``."""
    for idx, header in enumerate(headers):
        if _is_loss_number_header(header):
            return idx
    # fallback: первая колонка, начинающаяся с «убыток»
    for idx, header in enumerate(headers):
        if header.startswith("убыток"):
            return idx
    return None


def parse_prod_sheet(
    worksheet: Worksheet,
    *,
    group_id: str,
    source_file: str,
    source_sheet: str,
) -> tuple[pd.DataFrame, ProdSheetStats]:
    """Разбирает один лист в таблицу участников группы.

    Args:
        worksheet: лист openpyxl.
        group_id: идентификатор группы (обычно stem файла).
        source_file: имя файла.
        source_sheet: имя листа.

    Returns:
        Кортеж «таблица участников, статистика листа».
    """
    found = find_header_row(worksheet)
    if found is None:
        return _empty_frame(), ProdSheetStats(skipped_no_header=1)

    header_row, headers, _raw = found
    loss_col = find_loss_column(headers)
    if loss_col is None:
        return _empty_frame(), ProdSheetStats(skipped_no_header=1)

    specs = build_name_birth_specs(headers)
    max_col = len(headers)
    max_row = worksheet.max_row or header_row
    records: list[dict[str, object]] = []
    colored = 0
    loss_numbers_seen: set[str] = set()

    for row in range(header_row + 1, max_row + 1):
        if not row_has_color(worksheet, row, max_col):
            continue
        colored += 1
        values = [_cell_value(worksheet.cell(row=row, column=col)) for col in range(1, max_col + 1)]
        loss_number = _normalize_loss_number(values[loss_col] if loss_col < len(values) else None)
        if loss_number:
            loss_numbers_seen.add(loss_number)

        if specs:
            for spec in specs:
                name_parts = [
                    normalize_text(values[idx]) if idx < len(values) else ""
                    for idx in spec.name_cols
                ]
                full_name = " ".join(part for part in name_parts if part).strip()
                if not full_name:
                    continue
                birth_raw = (
                    values[spec.birth_col]
                    if spec.birth_col is not None and spec.birth_col < len(values)
                    else None
                )
                birth_date = pd.to_datetime(birth_raw, errors="coerce")
                if pd.isna(birth_date):
                    birth_date = pd.NaT
                else:
                    birth_date = birth_date.normalize()
                person_key = format_person_key(full_name, birth_date)
                if not person_key:
                    continue
                records.append(
                    {
                        "full_name": full_name,
                        "birth_date": birth_date,
                        "group_id": group_id,
                        "person_key": person_key,
                        "loss_number": loss_number,
                        "source_file": source_file,
                        "source_sheet": source_sheet,
                        "row_index": row,
                    }
                )
        elif loss_number:
            # Только номер убытка — для матча по LossNumber
            records.append(
                {
                    "full_name": "",
                    "birth_date": pd.NaT,
                    "group_id": group_id,
                    "person_key": "",
                    "loss_number": loss_number,
                    "source_file": source_file,
                    "source_sheet": source_sheet,
                    "row_index": row,
                }
            )

    frame = pd.DataFrame.from_records(records, columns=list(OUTPUT_COLUMNS)) if records else _empty_frame()
    if not frame.empty:
        frame = frame.drop_duplicates(
            subset=["group_id", "person_key", "loss_number"], keep="first"
        ).reset_index(drop=True)

    stats = ProdSheetStats(
        sheets_with_ubytok=1,
        colored_rows=colored,
        loss_numbers=len(loss_numbers_seen),
        person_keys=int(frame["person_key"].astype(bool).sum()) if not frame.empty else 0,
    )
    return frame, stats


def _empty_frame() -> pd.DataFrame:
    return pd.DataFrame(columns=list(OUTPUT_COLUMNS))


def read_prod_workbook(path: str | Path) -> tuple[pd.DataFrame, ProdSheetStats]:
    """Читает xlsx: все подходящие листы → одна группа = stem файла.

    Args:
        path: путь к файлу.

    Returns:
        Кортеж «таблица участников, агрегированная статистика по листам».
    """
    path = Path(path)
    group_id = path.stem
    frames: list[pd.DataFrame] = []
    total = ProdSheetStats()

    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as error:  # noqa: BLE001
        LOGGER.error("Не удалось открыть prod-xlsx %s: %s", path.name, error)
        return _empty_frame(), total

    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            state = getattr(worksheet, "sheet_state", "visible")
            if state == "hidden":
                continue
            frame, stats = parse_prod_sheet(
                worksheet,
                group_id=group_id,
                source_file=path.name,
                source_sheet=sheet_name,
            )
            total = ProdSheetStats(
                sheets_with_ubytok=total.sheets_with_ubytok + stats.sheets_with_ubytok,
                colored_rows=total.colored_rows + stats.colored_rows,
                loss_numbers=total.loss_numbers + stats.loss_numbers,
                person_keys=total.person_keys + stats.person_keys,
                skipped_no_header=total.skipped_no_header + stats.skipped_no_header,
            )
            if not frame.empty:
                frames.append(frame)
    finally:
        workbook.close()

    if not frames:
        return _empty_frame(), total

    combined = pd.concat(frames, ignore_index=True)
    combined["group_id"] = group_id
    combined = combined.drop_duplicates(
        subset=["group_id", "person_key", "loss_number"], keep="first"
    ).reset_index(drop=True)
    LOGGER.info(
        "Prod Excel %s: листов с Убыток=%d, цветных строк=%d, участников=%d",
        path.name,
        total.sheets_with_ubytok,
        total.colored_rows,
        len(combined),
    )
    return combined, total


def load_prod_beneficiaries(paths: Iterable[str | Path]) -> tuple[pd.DataFrame, ProdSheetStats]:
    """Читает список prod-xlsx в одну таблицу."""
    frames: list[pd.DataFrame] = []
    total = ProdSheetStats()
    for path in paths:
        frame, stats = read_prod_workbook(path)
        total = ProdSheetStats(
            sheets_with_ubytok=total.sheets_with_ubytok + stats.sheets_with_ubytok,
            colored_rows=total.colored_rows + stats.colored_rows,
            loss_numbers=total.loss_numbers + stats.loss_numbers,
            person_keys=total.person_keys + stats.person_keys,
            skipped_no_header=total.skipped_no_header + stats.skipped_no_header,
        )
        if not frame.empty:
            frames.append(frame)
    if not frames:
        return _empty_frame(), total
    return pd.concat(frames, ignore_index=True), total
