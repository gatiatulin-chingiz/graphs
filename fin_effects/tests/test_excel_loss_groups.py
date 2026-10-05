"""Тесты prod-парсера Excel (колонка «Убыток» + цветные строки)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from fin_effects.excel_loss_groups import (
    build_name_birth_specs,
    normalize_header,
    read_prod_workbook,
)
from fin_effects.preprocess import (
    attach_groups_by_loss_number,
    attach_groups_to_losses,
    collapse_losses_to_facts,
    normalize_losses,
)


def test_normalize_header_case_and_spaces() -> None:
    """Регистр и пробелы не влияют на сравнение заголовков."""
    assert normalize_header("Убыток") == "убыток"
    assert normalize_header("УБЫТОК") == "убыток"
    assert normalize_header("  Убыток  ") == "убыток"
    assert normalize_header("Убыток.\u00a0Код") == "убыток.код"
    assert normalize_header(None) == ""


def test_build_name_birth_specs_pairs_roles() -> None:
    """ФИО и ДР одной роли связываются по префиксу."""
    headers = [
        "убыток",
        "виновник полн наименование",
        "виновник дата рождения",
        "потерпевший полн наименование",
        "потерпевший дата рождения",
    ]
    specs = build_name_birth_specs(headers)
    assert len(specs) >= 2
    name_cols = {spec.name_cols[0] for spec in specs if len(spec.name_cols) == 1}
    assert 1 in name_cols and 3 in name_cols


def _write_prod_workbook(path: Path, *, upper_loss: bool = False) -> None:
    """Создаёт prod-like xlsx: шапка + цветная и обычная строки."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Группа"
    loss_header = "УБЫТОК" if upper_loss else "Убыток"
    headers = [
        loss_header,
        "Виновник Полн наименование",
        "Виновник Дата рождения",
        "Потерпевший Полн наименование",
        "Потерпевший Дата рождения",
    ]
    sheet.append(headers)
    # Цветная строка (в группе)
    sheet.append(
        [
            "LN-100",
            "Иванов Иван Иванович",
            "1980-01-15",
            "Петров Пётр Петрович",
            "1990-02-20",
        ]
    )
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    for col in range(1, 6):
        sheet.cell(row=2, column=col).fill = yellow

    # Строка только с цветным текстом
    sheet.append(
        [
            "LN-200",
            "Сидоров Сидор",
            "1975-03-10",
            "",
            "",
        ]
    )
    sheet.cell(row=3, column=2).font = Font(color="FF0000")

    # Некрашенная строка — не должна попасть
    sheet.append(
        [
            "LN-999",
            "Не должен попасть",
            "2000-01-01",
            "",
            "",
        ]
    )

    # Второй лист без «Убыток» — пропускается
    other = workbook.create_sheet("Сводка")
    other.append(["Группа", "Комментарий"])
    other.append(["1", "ignore"])

    workbook.save(path)


def test_read_prod_workbook_colored_rows_and_case(tmp_path: Path) -> None:
    """Цветные строки (заливка/шрифт) попадают в группу; регистр «УБЫТОК» ок."""
    path = tmp_path / "grp_demo.xlsx"
    _write_prod_workbook(path, upper_loss=True)

    frame, stats = read_prod_workbook(path)
    assert stats.sheets_with_ubytok == 1
    assert stats.colored_rows == 2
    assert set(frame["loss_number"]) == {"ln-100", "ln-200"}
    assert "не должен попасть" not in set(frame["full_name"])
    assert frame["group_id"].unique().tolist() == ["grp_demo"]
    assert frame["person_key"].astype(bool).any()
    keys = set(frame["person_key"])
    assert any("иванов" in key for key in keys)
    assert any("сидоров" in key for key in keys)


def test_read_prod_skips_sheet_without_ubytok(tmp_path: Path) -> None:
    """Лист без колонки Убыток не даёт участников."""
    path = tmp_path / "empty_style.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ФИО", "Дата рождения", "Группа"])
    sheet.append(["Иванов", "1980-01-01", "1"])
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    sheet.cell(row=2, column=1).fill = yellow
    workbook.save(path)

    frame, stats = read_prod_workbook(path)
    assert frame.empty
    assert stats.sheets_with_ubytok == 0


def test_attach_groups_by_loss_number(synthetic) -> None:
    """Матч по номеру убытка из Excel привязывает факты к группе."""
    normalized = normalize_losses(synthetic.losses)
    facts, _ = collapse_losses_to_facts(normalized)
    sample = facts.head(3).copy()
    beneficiaries = pd.DataFrame(
        {
            "group_id": ["excel_group"] * 3,
            "person_key": ["", "", ""],
            "loss_number": sample["LossNumber"].tolist(),
            "full_name": ["", "", ""],
            "birth_date": [pd.NaT, pd.NaT, pd.NaT],
        }
    )
    membership = attach_groups_by_loss_number(facts, beneficiaries)
    assert len(membership) == 3
    assert set(membership["group_id"]) == {"excel_group"}

    _, paired = attach_groups_to_losses(facts, beneficiaries)
    assert set(paired["group_id"]) == {"excel_group"}
    assert paired["LossNumber"].nunique() == 3


def test_attach_loss_number_strips_float_dot_zero() -> None:
    """SQL float LossNumber (123.0) матчится с Excel-строкой «123»."""
    from fin_effects.preprocess import normalize_losses
    from fin_effects.textnorm import normalize_loss_key

    assert normalize_loss_key(pd.Series([12345.0, "12345", " 12345.0 "])).tolist() == [
        "12345",
        "12345",
        "12345",
    ]

    facts = pd.DataFrame({"LossNumber": normalize_loss_key(pd.Series([1001.0, 1002.0]))})
    beneficiaries = pd.DataFrame(
        {
            "group_id": ["g", "g"],
            "loss_number": ["1001", "1002"],
            "person_key": ["", ""],
        }
    )
    membership = attach_groups_by_loss_number(facts, beneficiaries)
    assert set(membership["LossNumber"]) == {"1001", "1002"}


def test_attach_unions_person_key_and_loss_number(synthetic) -> None:
    """Union: пары по person_key и по loss_number объединяются."""
    normalized = normalize_losses(synthetic.losses)
    facts, _ = collapse_losses_to_facts(normalized)
    _, via_person = attach_groups_to_losses(facts, synthetic.beneficiaries)
    assert not via_person.empty

    extra_loss = facts["LossNumber"].iloc[0]
    # Убедимся, что номер не из уже сопоставленных (если попадёт — ок, union)
    extra = pd.DataFrame(
        {
            "group_id": ["extra_by_loss"],
            "person_key": [""],
            "loss_number": [extra_loss],
            "full_name": [""],
            "birth_date": [pd.NaT],
        }
    )
    combined = pd.concat([synthetic.beneficiaries, extra], ignore_index=True)
    _, paired = attach_groups_to_losses(facts, combined, match_mode="union")
    assert "extra_by_loss" in set(paired["group_id"])


def test_attach_auto_skips_person_key_when_loss_numbers_present(synthetic) -> None:
    """В auto при номерах убытков в Excel не тянем все убытки участника из КАСКО."""
    normalized = normalize_losses(synthetic.losses)
    facts, _ = collapse_losses_to_facts(normalized)
    sample = facts.head(5).copy()
    beneficiaries = pd.DataFrame(
        {
            "group_id": ["g1"] * 5,
            "person_key": synthetic.beneficiaries["person_key"].head(5).tolist(),
            "loss_number": sample["LossNumber"].tolist(),
            "full_name": [""] * 5,
            "birth_date": [pd.NaT] * 5,
        }
    )
    _, via_union = attach_groups_to_losses(facts, beneficiaries, match_mode="union")
    _, via_auto = attach_groups_to_losses(facts, beneficiaries, match_mode="auto")
    assert len(via_auto) <= len(via_union)
    assert via_auto["LossNumber"].nunique() == 5
